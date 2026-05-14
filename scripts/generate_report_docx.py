from __future__ import annotations

import re
from datetime import date
from pathlib import Path

import markdown
from bs4 import BeautifulSoup, NavigableString, Tag
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt


ROOT = Path(__file__).resolve().parents[1]
MD_PATH = ROOT / "报告.md"
DOCX_PATH = ROOT / "报告.docx"


def set_run_font(run, east_asia: str = "宋体", western: str = "Times New Roman", size: float | None = None,
                 bold: bool | None = None, italic: bool | None = None):
    run.font.name = western
    run._element.rPr.rFonts.set(qn("w:eastAsia"), east_asia)
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic


def set_paragraph_border_bottom(paragraph):
    p_pr = paragraph._p.get_or_add_pPr()
    p_bdr = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "8")
    bottom.set(qn("w:space"), "1")
    bottom.set(qn("w:color"), "000000")
    p_bdr.append(bottom)
    p_pr.append(p_bdr)


def set_page_setup(section):
    section.page_width = Cm(21)
    section.page_height = Cm(29.7)
    section.left_margin = Cm(2.5)
    section.right_margin = Cm(2)
    section.top_margin = Cm(2)
    section.bottom_margin = Cm(2)
    section.header_distance = Cm(1.5)
    section.footer_distance = Cm(1.75)


def exact_18(paragraph):
    fmt = paragraph.paragraph_format
    fmt.line_spacing_rule = WD_LINE_SPACING.EXACTLY
    fmt.line_spacing = Pt(18)


def add_field_run(paragraph, instruction: str):
    run = paragraph.add_run()
    fld_begin = OxmlElement("w:fldChar")
    fld_begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = instruction
    fld_sep = OxmlElement("w:fldChar")
    fld_sep.set(qn("w:fldCharType"), "separate")
    fld_end = OxmlElement("w:fldChar")
    fld_end.set(qn("w:fldCharType"), "end")
    run._r.append(fld_begin)
    run._r.append(instr)
    run._r.append(fld_sep)
    hint = paragraph.add_run("右键更新域")
    set_run_font(hint, size=12)
    run._r.append(fld_end)
    return hint


def set_update_fields_on_open(document: Document):
    settings = document.settings._element
    node = settings.find(qn("w:updateFields"))
    if node is None:
        node = OxmlElement("w:updateFields")
        settings.append(node)
    node.set(qn("w:val"), "true")


def restart_page_numbering(section, start: int = 1):
    sect_pr = section._sectPr
    pg_num = sect_pr.find(qn("w:pgNumType"))
    if pg_num is None:
        pg_num = OxmlElement("w:pgNumType")
        sect_pr.append(pg_num)
    pg_num.set(qn("w:start"), str(start))


def build_styles(document: Document):
    normal = document.styles["Normal"]
    normal.font.name = "Times New Roman"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
    normal.font.size = Pt(12)
    pf = normal.paragraph_format
    pf.first_line_indent = Pt(24)
    pf.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    pf.space_before = Pt(0)
    pf.space_after = Pt(0)
    pf.line_spacing_rule = WD_LINE_SPACING.EXACTLY
    pf.line_spacing = Pt(18)

    for style_name in ["Heading 1", "Heading 2", "Heading 3"]:
        style = document.styles[style_name]
        style.font.name = "Times New Roman"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "黑体")
        style.font.bold = False
        style.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
        style.paragraph_format.space_before = Pt(12)
        style.paragraph_format.space_after = Pt(12)
        style.paragraph_format.line_spacing_rule = WD_LINE_SPACING.EXACTLY
        style.paragraph_format.line_spacing = Pt(18)
        style.paragraph_format.first_line_indent = Pt(0)
    document.styles["Heading 1"].font.size = Pt(14)
    document.styles["Heading 2"].font.size = Pt(12)
    document.styles["Heading 3"].font.size = Pt(12)

    for toc_name, indent in [("TOC 1", 0), ("TOC 2", 24), ("TOC 3", 48)]:
        if toc_name not in document.styles:
            style = document.styles.add_style(toc_name, WD_STYLE_TYPE.PARAGRAPH)
        else:
            style = document.styles[toc_name]
        style.font.name = "Times New Roman"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
        style.font.size = Pt(12)
        style.paragraph_format.left_indent = Pt(indent)
        style.paragraph_format.space_before = Pt(0)
        style.paragraph_format.space_after = Pt(0)
        style.paragraph_format.line_spacing_rule = WD_LINE_SPACING.EXACTLY
        style.paragraph_format.line_spacing = Pt(18)
        style.paragraph_format.first_line_indent = Pt(0)

    def ensure(name):
        if name not in document.styles:
            return document.styles.add_style(name, WD_STYLE_TYPE.PARAGRAPH)
        return document.styles[name]

    styles = {}

    styles["cover_university"] = ensure("CoverUniversity")
    styles["cover_university"].font.name = "Times New Roman"
    styles["cover_university"]._element.rPr.rFonts.set(qn("w:eastAsia"), "黑体")
    styles["cover_university"].font.size = Pt(22)
    styles["cover_university"].font.bold = False
    styles["cover_university"].paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    styles["cover_university"].paragraph_format.space_before = Pt(0)
    styles["cover_university"].paragraph_format.space_after = Pt(0)
    styles["cover_university"].paragraph_format.line_spacing_rule = WD_LINE_SPACING.EXACTLY
    styles["cover_university"].paragraph_format.line_spacing = Pt(18)

    styles["cover_report"] = ensure("CoverReport")
    styles["cover_report"].font.name = "Times New Roman"
    styles["cover_report"]._element.rPr.rFonts.set(qn("w:eastAsia"), "黑体")
    styles["cover_report"].font.size = Pt(22)
    styles["cover_report"].paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    styles["cover_report"].paragraph_format.line_spacing_rule = WD_LINE_SPACING.EXACTLY
    styles["cover_report"].paragraph_format.line_spacing = Pt(18)

    styles["cover_item"] = ensure("CoverItem")
    styles["cover_item"].font.name = "Times New Roman"
    styles["cover_item"]._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
    styles["cover_item"].font.size = Pt(14)
    styles["cover_item"].paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
    styles["cover_item"].paragraph_format.line_spacing_rule = WD_LINE_SPACING.EXACTLY
    styles["cover_item"].paragraph_format.line_spacing = Pt(18)
    styles["cover_item"].paragraph_format.first_line_indent = Pt(0)

    styles["toc_title"] = ensure("TOCTitleCN")
    styles["toc_title"].font.name = "Times New Roman"
    styles["toc_title"]._element.rPr.rFonts.set(qn("w:eastAsia"), "黑体")
    styles["toc_title"].font.size = Pt(14)
    styles["toc_title"].paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    styles["toc_title"].paragraph_format.space_before = Pt(0)
    styles["toc_title"].paragraph_format.space_after = Pt(0)
    styles["toc_title"].paragraph_format.line_spacing_rule = WD_LINE_SPACING.EXACTLY
    styles["toc_title"].paragraph_format.line_spacing = Pt(18)

    styles["abstract_title"] = ensure("AbstractTitleCN")
    styles["abstract_title"].font.name = "Times New Roman"
    styles["abstract_title"]._element.rPr.rFonts.set(qn("w:eastAsia"), "黑体")
    styles["abstract_title"].font.size = Pt(14)
    styles["abstract_title"].paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    styles["abstract_title"].paragraph_format.space_before = Pt(12)
    styles["abstract_title"].paragraph_format.space_after = Pt(12)
    styles["abstract_title"].paragraph_format.line_spacing_rule = WD_LINE_SPACING.EXACTLY
    styles["abstract_title"].paragraph_format.line_spacing = Pt(18)

    styles["caption"] = ensure("FigureCaptionCN")
    styles["caption"].font.name = "Times New Roman"
    styles["caption"]._element.rPr.rFonts.set(qn("w:eastAsia"), "黑体")
    styles["caption"].font.size = Pt(10.5)
    styles["caption"].paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    styles["caption"].paragraph_format.space_before = Pt(6)
    styles["caption"].paragraph_format.space_after = Pt(6)
    styles["caption"].paragraph_format.line_spacing_rule = WD_LINE_SPACING.EXACTLY
    styles["caption"].paragraph_format.line_spacing = Pt(18)
    styles["caption"].paragraph_format.first_line_indent = Pt(0)

    styles["body"] = normal
    return styles


def add_cover(document: Document, styles, report_title: str):
    for _ in range(6):
        p = document.add_paragraph(style=styles["cover_item"])
        exact_18(p)

    p = document.add_paragraph("河 北 农 业 大 学", style=styles["cover_university"])
    exact_18(p)
    p = document.add_paragraph("软件工程课程设计报告", style=styles["cover_report"])
    exact_18(p)

    for _ in range(4):
        p = document.add_paragraph(style=styles["cover_item"])
        exact_18(p)

    p = document.add_paragraph(style=styles["cover_item"])
    exact_18(p)
    r = p.add_run("题    目：")
    set_run_font(r, east_asia="宋体", size=14)
    r = p.add_run(report_title)
    set_run_font(r, east_asia="宋体", size=14)

    for label, value in [
        ("组    名：", "____________________________"),
    ]:
        p = document.add_paragraph(style=styles["cover_item"])
        exact_18(p)
        r = p.add_run(label)
        set_run_font(r, east_asia="宋体", size=14)
        r = p.add_run(value)
        set_run_font(r, east_asia="宋体", size=14)

    for _ in range(8):
        p = document.add_paragraph(style=styles["cover_item"])
        exact_18(p)

    today = date.today()
    p = document.add_paragraph(style=styles["cover_item"])
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    exact_18(p)
    r = p.add_run(f"{today.year} 年 {today.month} 月 {today.day} 日")
    set_run_font(r, east_asia="宋体", size=14)


def add_toc(document: Document, styles):
    p = document.add_paragraph("目  录", style=styles["toc_title"])
    exact_18(p)
    p = document.add_paragraph()
    p.paragraph_format.first_line_indent = Pt(0)
    exact_18(p)
    add_field_run(p, 'TOC \\o "1-3" \\h \\z \\u')


def normalize_heading(text: str, level: int) -> str:
    text = re.sub(r"\s+", " ", text.strip())
    if level == 1:
        text = re.sub(r"^(\d+)\.\s*", r"\1 ", text)
    return text


def add_inline_runs(paragraph, node):
    if isinstance(node, NavigableString):
        text = str(node)
        if text:
            run = paragraph.add_run(text)
            set_run_font(run, size=12)
        return

    if not isinstance(node, Tag):
        return

    if node.name == "strong":
        run = paragraph.add_run(node.get_text())
        set_run_font(run, size=12, bold=True)
        return
    if node.name == "em":
        run = paragraph.add_run(node.get_text())
        set_run_font(run, size=12, italic=True)
        return
    if node.name == "code":
        run = paragraph.add_run(node.get_text())
        set_run_font(run, east_asia="等线", western="Consolas", size=12)
        return

    for child in node.children:
        add_inline_runs(paragraph, child)


def add_body_paragraph(document: Document, node, styles):
    p = document.add_paragraph(style=styles["body"])
    exact_18(p)
    for child in node.children:
        add_inline_runs(p, child)
    return p


def add_list(document: Document, list_tag: Tag, styles):
    ordered = list_tag.name == "ol"
    for i, li in enumerate(list_tag.find_all("li", recursive=False), start=1):
        p = document.add_paragraph(style=styles["body"])
        exact_18(p)
        prefix = f"{i}. " if ordered else "• "
        run = p.add_run(prefix)
        set_run_font(run, size=12)
        for child in li.children:
            if isinstance(child, Tag) and child.name in {"ol", "ul"}:
                add_list(document, child, styles)
            else:
                add_inline_runs(p, child)


def add_image(document: Document, src: str, alt: str, styles, figure_no: int):
    image_path = (ROOT / src).resolve()
    p = document.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.keep_with_next = True
    exact_18(p)
    run = p.add_run()
    run.add_picture(str(image_path), width=Cm(15.5))

    caption = document.add_paragraph(style=styles["caption"])
    caption.paragraph_format.keep_together = True
    exact_18(caption)
    text = f"图{figure_no} {alt}"
    run = caption.add_run(text)
    set_run_font(run, east_asia="黑体", size=10.5)


def render_body(document: Document, html: str, styles):
    soup = BeautifulSoup(html, "html.parser")
    body_nodes = soup.contents
    figure_no = 1
    first_title_skipped = False

    for node in body_nodes:
        if isinstance(node, NavigableString):
            continue
        if not isinstance(node, Tag):
            continue

        if node.name == "h1" and not first_title_skipped:
            first_title_skipped = True
            continue

        if node.name == "h2":
            text = normalize_heading(node.get_text(), 0)
            if text == "摘要":
                p = document.add_paragraph(text, style=styles["abstract_title"])
            else:
                p = document.add_paragraph(normalize_heading(text, 1), style="Heading 1")
            exact_18(p)
            continue

        if node.name == "h3":
            p = document.add_paragraph(normalize_heading(node.get_text(), 2), style="Heading 2")
            exact_18(p)
            continue

        if node.name == "h4":
            p = document.add_paragraph(normalize_heading(node.get_text(), 3), style="Heading 3")
            exact_18(p)
            continue

        if node.name == "p":
            imgs = node.find_all("img", recursive=False)
            if len(imgs) == 1 and not node.get_text(strip=True):
                img = imgs[0]
                alt = img.get("alt", "插图")
                add_image(document, img.get("src"), alt, styles, figure_no)
                figure_no += 1
            else:
                add_body_paragraph(document, node, styles)
            continue

        if node.name in {"ol", "ul"}:
            add_list(document, node, styles)
            continue

        if node.name == "hr":
            continue


def add_footer_page_number(section):
    section.footer.is_linked_to_previous = False
    footer = section.footer
    p = footer.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    p.paragraph_format.first_line_indent = Pt(0)
    exact_18(p)
    run = p.add_run()
    set_run_font(run, east_asia="Times New Roman", western="Times New Roman", size=9)
    fld_begin = OxmlElement("w:fldChar")
    fld_begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    fld_end = OxmlElement("w:fldChar")
    fld_end.set(qn("w:fldCharType"), "end")
    run._r.append(fld_begin)
    run._r.append(instr)
    run._r.append(fld_end)


def main():
    md_text = MD_PATH.read_text(encoding="utf-8")
    title_match = re.search(r"^#\s+(.+)$", md_text, flags=re.M)
    report_title = title_match.group(1).strip() if title_match else "软件工程课程设计报告"
    html = markdown.markdown(md_text, extensions=["extra"])

    document = Document()
    set_update_fields_on_open(document)
    styles = build_styles(document)

    for section in document.sections:
        set_page_setup(section)

    add_cover(document, styles, report_title)

    toc_section = document.add_section(WD_SECTION.NEW_PAGE)
    set_page_setup(toc_section)
    add_toc(document, styles)

    body_section = document.add_section(WD_SECTION.NEW_PAGE)
    set_page_setup(body_section)
    restart_page_numbering(body_section, 1)
    add_footer_page_number(body_section)
    render_body(document, html, styles)

    document.save(str(DOCX_PATH))


if __name__ == "__main__":
    main()
